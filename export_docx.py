"""
Generates a branded Treasury Aesthetics Word document (.docx) from a treatment plan JSON.
"""

import io
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# Brand colours
CHARCOAL = RGBColor(0x1A, 0x1A, 0x1A)
GOLD     = RGBColor(0xC4, 0x95, 0x4A)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
GREY     = RGBColor(0x6B, 0x6B, 0x6B)
LIGHT    = RGBColor(0xF5, 0xF2, 0xEE)


def _set_cell_bg(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)


def _hr(doc, color_hex='C4954A', thickness=6):
    """Insert a thin horizontal rule paragraph."""
    p    = doc.add_paragraph()
    pPr  = p._p.get_or_add_pPr()
    pb   = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    str(thickness))
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), color_hex)
    pb.append(bot)
    pPr.append(pb)
    p.paragraph_format.space_after  = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    return p


def generate_plan_docx(plan: dict) -> bytes:
    """Return .docx bytes for the given plan dict."""
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # ── Header banner (1-cell table) ──────────────────────────────────────────
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.cell(0, 0)
    _set_cell_bg(cell, '1A1A1A')
    cell.width = Inches(6.5)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run('TREASURY AESTHETICS')
    run.bold      = True
    run.font.size = Pt(16)
    run.font.color.rgb = GOLD
    run.font.name      = 'Georgia'

    p2 = cell.add_paragraph('Dr. Jason Latsky, MD  ·  Toronto, Ontario')
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after  = Pt(10)
    r2 = p2.runs[0]
    r2.font.size  = Pt(9)
    r2.font.color.rgb = RGBColor(0xB0, 0xA8, 0x9E)
    r2.font.name  = 'Calibri'

    doc.add_paragraph()

    # ── Plan title ────────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_after = Pt(2)
    tr = title_p.add_run(plan.get('title', 'Treatment Plan'))
    tr.bold           = True
    tr.font.size      = Pt(18)
    tr.font.color.rgb = CHARCOAL
    tr.font.name      = 'Georgia'

    date_p = doc.add_paragraph(f'Prepared: {date.today().strftime("%B %d, %Y")}')
    date_p.paragraph_format.space_after = Pt(4)
    date_p.runs[0].font.size      = Pt(9)
    date_p.runs[0].font.color.rgb = GREY
    date_p.runs[0].font.name      = 'Calibri'

    _hr(doc)
    doc.add_paragraph()

    # ── Sections ──────────────────────────────────────────────────────────────
    for section in plan.get('sections', []):
        items = section.get('items', [])
        if not items:
            continue

        # Section heading
        sh = doc.add_paragraph()
        sh.paragraph_format.space_before = Pt(10)
        sh.paragraph_format.space_after  = Pt(4)
        sr = sh.add_run(section.get('name', '').upper())
        sr.bold           = True
        sr.font.size      = Pt(9)
        sr.font.color.rgb = GOLD
        sr.font.name      = 'Calibri'
        sr.font.all_caps  = True

        for item in items:
            checked = item.get('checked', True)
            name    = item.get('name', '')
            detail  = item.get('detail', '')

            marker = '✓' if checked else '○'
            ip = doc.add_paragraph(style='Normal')
            ip.paragraph_format.left_indent  = Cm(0.5)
            ip.paragraph_format.space_before = Pt(3)
            ip.paragraph_format.space_after  = Pt(1)

            mr = ip.add_run(f'{marker}  {name}')
            mr.bold           = checked
            mr.font.size      = Pt(10.5)
            mr.font.color.rgb = CHARCOAL if checked else GREY
            mr.font.name      = 'Calibri'

            if detail:
                dp = doc.add_paragraph(detail)
                dp.paragraph_format.left_indent  = Cm(1.2)
                dp.paragraph_format.space_before = Pt(0)
                dp.paragraph_format.space_after  = Pt(2)
                dp.runs[0].font.size      = Pt(9)
                dp.runs[0].font.color.rgb = GREY
                dp.runs[0].font.name      = 'Calibri'
                dp.runs[0].italic         = True

        doc.add_paragraph()

    # ── Footer ────────────────────────────────────────────────────────────────
    _hr(doc, color_hex='2A2520', thickness=4)
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(8)
    fr = fp.add_run(
        'This plan was prepared exclusively for you by the Treasury Aesthetics team.\n'
        'aesthetics@treasuryhealth.ca  ·  treasuryhealth.ca  ·  Toronto, Ontario'
    )
    fr.font.size      = Pt(8)
    fr.font.color.rgb = GREY
    fr.font.name      = 'Calibri'
    fr.italic         = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
